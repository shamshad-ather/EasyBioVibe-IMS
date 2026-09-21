import sqlite3
from io import BytesIO
from openpyxl import Workbook, load_workbook
from flask import Blueprint, request, jsonify, session, make_response
from app.database import get_db, trigger_backup
from app.utils import login_required, admin_required, generate_semantic_id

equipment_bp = Blueprint('equipment', __name__)

@equipment_bp.route('/api/equipment', methods=['GET', 'POST'])
@login_required
def handle_equipment():
    conn = get_db()
    c = conn.cursor()
    if request.method == 'POST':
        if session.get('role') != 'Admin':
            conn.close()
            return jsonify({"status": "error", "message": "Admin privileges required"}), 403
        data = request.get_json()
        try:
            if data.get('id'):
                c.execute("""UPDATE Equipment_Master SET name=?, make=?, model=?, serial_number=?, department_id=?, faculty_in_charge=?, installation_date=?, location_room=?, status=? WHERE id=?""",
                          (data.get('name'), data.get('make', ''), data.get('model', ''), data.get('serial_number', ''), data.get('department_id'), data.get('faculty_in_charge', ''), data.get('installation_date', ''), data.get('location_room', ''), data.get('status', 'Active'), data.get('id')))
            else:
                c.execute("SELECT MAX(id) FROM Equipment_Master")
                max_id = c.fetchone()[0] or 0
                code = data.get('equip_code') or generate_semantic_id('EQ', max_id + 1, conn)
                c.execute("""INSERT INTO Equipment_Master (equip_code, name, make, model, serial_number, department_id, faculty_in_charge, installation_date, location_room, status) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                          (code, data.get('name'), data.get('make', ''), data.get('model', ''), data.get('serial_number', ''), data.get('department_id'), data.get('faculty_in_charge', ''), data.get('installation_date', ''), data.get('location_room', ''), data.get('status', 'Active')))
            conn.commit()
            status, msg = "success", "Asset saved successfully"
        except sqlite3.IntegrityError as e:
            status, msg = "error", f"Database constraint violation: {str(e)}"
        except Exception as e:
            status, msg = "error", str(e)
        conn.close()
        if status == "success": trigger_backup()
        return jsonify({"status": status, "message": msg})
        
    c.execute("SELECT * FROM Equipment_Master ORDER BY id DESC")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return jsonify(rows)

@equipment_bp.route('/api/equipment/events', methods=['GET', 'POST'])
@login_required
def handle_equipment_events():
    conn = get_db()
    c = conn.cursor()
    if request.method == 'POST':
        data = request.get_json()
        if not data.get('equip_id'): 
            conn.close()
            return jsonify({"status": "error", "message": "Equipment ID is required."}), 400
            
        if data.get('id'): 
            c.execute("""UPDATE Equipment_Events SET event_type=?, event_date=?, performed_by=?, pass_fail_status=?, remarks=? WHERE id=?""", 
                      (data.get('event_type'), data.get('event_date'), data.get('performed_by', ''), data.get('pass_fail_status', ''), data.get('remarks', ''), data.get('id')))
        else: 
            c.execute("""INSERT INTO Equipment_Events (equip_id, event_type, event_date, performed_by, pass_fail_status, remarks) VALUES (?, ?, ?, ?, ?, ?)""", 
                      (data.get('equip_id'), data.get('event_type'), data.get('event_date'), data.get('performed_by', ''), data.get('pass_fail_status', ''), data.get('remarks', '')))
        conn.commit()
        conn.close()
        trigger_backup()
        return jsonify({"status": "success"})
        
    c.execute("SELECT * FROM Equipment_Events ORDER BY event_date DESC, id DESC")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return jsonify(rows)

@equipment_bp.route('/api/equipment/template', methods=['GET'])
@admin_required
def equipment_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Equipment Upload"
    headers = ['Name', 'Make', 'Model', 'SerialNumber', 'DepartmentID', 'FacultyInCharge', 'InstallationDate', 'LocationRoom', 'Status']
    ws.append(headers)
    ws.append(['PCR Machine', 'Bio-Rad', 'T100', 'SN123456', 1, 'Dr. Smith', '2025-01-01', 'Room 101', 'Active'])
    
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    response = make_response(out.read())
    response.headers['Content-Disposition'] = 'attachment; filename=Equipment_Template.xlsx'
    response.headers['Content-type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response

@equipment_bp.route('/api/equipment/upload', methods=['POST'])
@admin_required
def upload_equipment():
    file = request.files.get('file')
    if not file: return jsonify({"status": "error", "message": "No file uploaded"}), 400
    try:
        wb = load_workbook(file, data_only=True)
        ws = wb.active
        rows = list(ws.rows)
        if len(rows) < 2: return jsonify({"status": "error", "message": "File is empty"}), 400
        headers = [str(cell.value).strip() if cell.value else f"Col{i}" for i, cell in enumerate(rows[0])]
        parsed_data = []
        for row in rows[1:]:
            row_data = {headers[i]: cell.value for i, cell in enumerate(row)}
            name = str(row_data.get('Name', '') or '').strip()
            if not name or name == 'None': continue
            
            parsed_data.append({
                "Name": name,
                "Make": str(row_data.get('Make', '') or '').strip(),
                "Model": str(row_data.get('Model', '') or '').strip(),
                "SerialNumber": str(row_data.get('SerialNumber', '') or '').strip(),
                "DepartmentID": row_data.get('DepartmentID'),
                "FacultyInCharge": str(row_data.get('FacultyInCharge', '') or '').strip(),
                "InstallationDate": str(row_data.get('InstallationDate', '') or '').strip(),
                "LocationRoom": str(row_data.get('LocationRoom', '') or '').strip(),
                "Status": str(row_data.get('Status', 'Active') or 'Active').strip()
            })
        return jsonify({"status": "success", "data": parsed_data})
    except Exception as e: return jsonify({"status": "error", "message": f"Failed to parse Excel: {str(e)}"}), 400

@equipment_bp.route('/api/equipment/bulk', methods=['POST'])
@admin_required
def bulk_equipment():
    data = request.get_json()
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT MAX(id) FROM Equipment_Master")
    max_id = c.fetchone()[0] or 0
    added_count = 0
    for row in data.get('rows', []):
        max_id += 1
        try:
            code = generate_semantic_id('EQ', max_id, conn)
            c.execute("""INSERT INTO Equipment_Master (equip_code, name, make, model, serial_number, department_id, faculty_in_charge, installation_date, location_room, status) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                      (code, row.get('Name'), row.get('Make', ''), row.get('Model', ''), row.get('SerialNumber', ''), row.get('DepartmentID'), row.get('FacultyInCharge', ''), row.get('InstallationDate', ''), row.get('LocationRoom', ''), row.get('Status', 'Active')))
            added_count += 1
        except sqlite3.IntegrityError: pass
    conn.commit()
    conn.close()
    trigger_backup()
    return jsonify({"status": "success", "message": f"Import complete: {added_count} added."})