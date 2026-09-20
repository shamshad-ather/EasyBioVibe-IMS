import sqlite3
import difflib
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from flask import Blueprint, request, jsonify, make_response, session
from app.database import get_db, trigger_backup
from app.utils import login_required, admin_required, generate_semantic_id

inventory_bp = Blueprint('inventory', __name__)

@inventory_bp.route('/api/inventory', methods=['GET', 'POST'])
@login_required
def handle_inventory():
    conn = get_db()
    c = conn.cursor()
    if request.method == 'POST':
        if session.get('role') != 'Admin': conn.close(); return jsonify({"status": "error", "message": "Admin privileges required"}), 403
        data = request.get_json()
        try:
            if data.get('id'):
                c.execute("""UPDATE Inventory_Master SET material_name=?, make=?, model=?, category=?, alert_threshold=?, base_unit=?, pack_qty=? WHERE id=?""",
                          (data['material_name'], data.get('make', ''), data.get('model', ''), data.get('category', 'Other'), float(data.get('alert_threshold', 15)), data.get('base_unit', 'Nos'), float(data.get('pack_qty', 1)), data.get('id')))
            else:
                c.execute("SELECT MAX(id) FROM Inventory_Master")
                max_id = c.fetchone()[0] or 0
                code = data.get('item_code') or generate_semantic_id('INV', max_id + 1, conn)
                c.execute("""INSERT INTO Inventory_Master (item_code, material_name, make, model, category, alert_threshold, base_unit, pack_qty) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                          (code, data['material_name'], data.get('make', ''), data.get('model', ''), data.get('category', 'Other'), float(data.get('alert_threshold', 15)), data.get('base_unit', 'Nos'), float(data.get('pack_qty', 1))))
            conn.commit(); status = "success"
        except sqlite3.IntegrityError: status = "error"
        conn.close()
        if status == "success": trigger_backup()
        return jsonify({"status": status})
    c.execute("SELECT * FROM Inventory_Master")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return jsonify(rows)

@inventory_bp.route('/api/inventory/upload', methods=['POST'])
@admin_required
def upload_inventory():
    file = request.files.get('file')
    if not file: return jsonify({"status": "error", "message": "No file uploaded"}), 400
    try:
        wb = load_workbook(file, data_only=True)
        ws = wb.active
        rows = list(ws.rows)
        if len(rows) < 2: return jsonify({"status": "error", "message": "File is empty"}), 400
        headers = [str(cell.value).strip() if cell.value else f"Col{i}" for i, cell in enumerate(rows[0])]
        conn = get_db(); c = conn.cursor()
        c.execute("SELECT id, material_name FROM Inventory_Master"); existing_items = [dict(r) for r in c.fetchall()]
        conn.close()
        parsed_data = []
        for row in rows[1:]:
            row_data = {headers[i]: cell.value for i, cell in enumerate(row)}
            new_name = str(row_data.get('MaterialName', '') or '').strip()
            if not new_name or new_name == 'None': continue
            best_match, best_score = None, 0.0
            for ext in existing_items:
                score = difflib.SequenceMatcher(None, new_name.lower(), ext['material_name'].lower()).ratio()
                if score > best_score: best_score, best_match = score, ext
            parsed_data.append({"MaterialName": new_name, "MaterialType": str(row_data.get('MaterialType', 'Other') or 'Other').strip(), "Make": str(row_data.get('Make', '') or '').strip(), "Model": str(row_data.get('Model', '') or '').strip(), "PackQty": float(row_data.get('PackQty') or 1), "Unit": str(row_data.get('Unit', 'Nos') or 'Nos').strip(), "AlertThreshold": float(row_data.get('AlertThreshold') or 15), "Description": str(row_data.get('Description', '') or '').strip(), "MatchScore": round(best_score * 100, 2), "MatchedID": best_match['id'] if best_match else None, "MatchedName": best_match['material_name'] if best_match else None})
        return jsonify({"status": "success", "data": parsed_data})
    except Exception as e: return jsonify({"status": "error", "message": f"Failed to parse Excel: {str(e)}"}), 400

@inventory_bp.route('/api/inventory/bulk', methods=['POST'])
@admin_required
def bulk_inventory():
    data = request.get_json()
    conn = get_db(); c = conn.cursor()
    c.execute("SELECT MAX(id) FROM Inventory_Master")
    max_id = c.fetchone()[0] or 0
    added_count, merged_count = 0, 0
    for row in data.get('rows', []):
        if row.get('MergeWithID'): merged_count += 1; continue
        max_id += 1
        try:
            code = generate_semantic_id('INV', max_id, conn)
            c.execute("""INSERT INTO Inventory_Master (item_code, material_name, make, model, category, alert_threshold, base_unit, pack_qty) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (code, row['MaterialName'], row.get('Make', ''), row.get('Model', ''), row.get('MaterialType', 'Other'), float(row.get('AlertThreshold', 15)), row.get('Unit', 'Nos'), float(row.get('PackQty', 1))))
            added_count += 1
        except sqlite3.IntegrityError: pass
    conn.commit(); conn.close(); trigger_backup()
    return jsonify({"status": "success", "message": f"Import complete: {added_count} added, {merged_count} skipped."})

@inventory_bp.route('/api/inventory/template', methods=['GET'])
@admin_required
def inventory_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory Upload"
    headers = ['MaterialName', 'MaterialType', 'Make', 'Model', 'PackQty', 'Unit', 'AlertThreshold', 'Description']
    ws.append(headers)
    ws.append(['Taq Polymerase', 'Reagent', 'ThermoFisher', '201-X', 500, 'Rxns', 15, 'Standard PCR enzyme'])
    units = ['Nos','uL','mL','L','ug','mg','g','kg','Box','Pack','Bottle','Rxns','Tube','Vial','Kit']
    dv_unit = DataValidation(type="list", formula1='"' + ','.join(units) + '"', allow_blank=True)
    ws.add_data_validation(dv_unit); dv_unit.add("F2:F1048576")
    mat_types = ['Reagent','Chemical','Plasticware','Glassware','Consumable','Kit','Buffer','Other']
    dv_type = DataValidation(type="list", formula1='"' + ','.join(mat_types) + '"', allow_blank=True)
    ws.add_data_validation(dv_type); dv_type.add("B2:B1048576")
    output = BytesIO()
    wb.save(output); output.seek(0)
    response = make_response(output.read())
    response.headers["Content-Disposition"] = "attachment; filename=EasyBio_Inventory_Template.xlsx"
    response.headers["Content-type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response