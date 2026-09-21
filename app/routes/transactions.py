import datetime
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from flask import Blueprint, request, jsonify, session, make_response
from app.database import get_db, trigger_backup
from app.utils import login_required, admin_required, get_unit_multiplier, generate_semantic_id

transactions_bp = Blueprint('transactions', __name__)

@transactions_bp.route('/api/batches', methods=['GET', 'POST'])
@login_required
def handle_batches():
    conn = get_db()
    c = conn.cursor()
    if request.method == 'POST':
        data = request.get_json()
        if not data.get('study_id'): conn.close(); return jsonify({"status": "error", "message": "Study is strictly required."}), 400
        if data.get('id'):
            c.execute("""UPDATE Physical_Batches SET inventory_id=?, po_number=?, lot_number=?, expiry_date=?, date_first_used=?, quantity_received=?, current_quantity=?, unit=?, department_id=?, study_id=?, vendor_id=?, status=?, remarks=? WHERE id=?""",
                      (data.get('inventory_id'), data.get('po_number', ''), data.get('lot_number', ''), data.get('expiry_date', ''), data.get('date_first_used', ''), data.get('quantity_received', 0), data.get('current_quantity', 0), data.get('unit', 'Nos'), int(data.get('department_id')) if data.get('department_id') else None, int(data.get('study_id')), int(data.get('vendor_id')) if data.get('vendor_id') else None, data.get('status', 'Active'), data.get('remarks', ''), data.get('id')))
        else:
            c.execute("SELECT MAX(id) FROM Physical_Batches")
            max_id = c.fetchone()[0] or 0
            code = generate_semantic_id('BAT', max_id + 1, conn)
            c.execute("""INSERT INTO Physical_Batches (batch_code, inventory_id, po_number, lot_number, expiry_date, date_first_used, quantity_received, current_quantity, unit, department_id, study_id, vendor_id, status, remarks) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                      (code, data.get('inventory_id'), data.get('po_number', ''), data.get('lot_number', ''), data.get('expiry_date', ''), data.get('date_first_used', ''), data.get('quantity_received', 0), data.get('current_quantity', 0), data.get('unit', 'Nos'), int(data.get('department_id')) if data.get('department_id') else None, int(data.get('study_id')), int(data.get('vendor_id')) if data.get('vendor_id') else None, data.get('status', 'Active'), data.get('remarks', '')))
        conn.commit(); conn.close(); trigger_backup()
        return jsonify({"status": "success"})
    c.execute("SELECT * FROM Physical_Batches ORDER BY id DESC")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return jsonify(rows)

@transactions_bp.route('/api/usage', methods=['GET', 'POST'])
@login_required
def handle_usage():
    conn = get_db()
    c = conn.cursor()
    
    if request.method == 'POST':
        data = request.get_json()
        
        try:
            qty_used = float(data.get('quantity_used', 0))
            if qty_used <= 0:
                conn.close()
                return jsonify({"status": "error", "message": "Quantity used must be greater than zero."}), 400
        except (ValueError, TypeError):
            conn.close()
            return jsonify({"status": "error", "message": "Invalid quantity specified."}), 400
        
        c.execute("SELECT current_quantity, unit, batch_code FROM Physical_Batches WHERE id = ?", (data.get('batch_id'),))
        batch = c.fetchone()
        if not batch: 
            conn.close()
            return jsonify({"status": "error", "message": "Batch not found"}), 400
            
        used_mult, used_fam = get_unit_multiplier(data.get('unit', ''))
        batch_mult, batch_fam = get_unit_multiplier(batch['unit'])
        
        if used_fam != batch_fam: 
            conn.close()
            return jsonify({"status": "error", "message": f"Unit mismatch: cannot convert {data.get('unit', '')} to {batch['unit']}."}), 400
            
        norm_qty = (qty_used * used_mult) / batch_mult
        
        if float(batch['current_quantity']) < norm_qty: 
            conn.close()
            return jsonify({"status": "error", "message": f"Insufficient stock. Available: {batch['current_quantity']} {batch['unit']}, Requested: {norm_qty} {batch['unit']}"}), 400
            
        new_qty = round(max(0.0, float(batch['current_quantity']) - norm_qty), 4)
        c.execute("UPDATE Physical_Batches SET current_quantity = ?, status = ? WHERE id = ? AND current_quantity >= ?", 
                  (new_qty, 'Active' if new_qty > 0 else 'Depleted', data.get('batch_id'), norm_qty))
        if c.rowcount == 0:
            conn.rollback()
            conn.close()
            return jsonify({"status": "error", "message": "Stock update conflict. Please try again."}), 409
        
        user_id = str(data.get('user_id', '') or '')
        user_name = session.get('user', '')
        if user_id:
            c.execute("SELECT username FROM Users WHERE id = ? OR username = ?", (user_id, user_id))
            u_row = c.fetchone()
            if u_row: user_name = u_row['username']

        dept_id = str(data.get('department_id', '') or '')
        dept_name = ''
        if dept_id:
            c.execute("SELECT name FROM Departments WHERE id = ? OR name = ?", (dept_id, dept_id))
            d_row = c.fetchone()
            if d_row: dept_name = d_row['name']

        fac_id = str(data.get('pi_user_id', data.get('faculty_id', '')) or '')
        fac_name = ''
        if fac_id:
            c.execute("SELECT username FROM Users WHERE id = ? OR username = ?", (fac_id, fac_id))
            f_row = c.fetchone()
            if f_row: fac_name = f_row['username']

        std_id = str(data.get('study_id', '') or '')
        std_name = ''
        if std_id:
            c.execute("SELECT name FROM Studies WHERE id = ? OR name = ?", (std_id, std_id))
            s_row = c.fetchone()
            if s_row: std_name = s_row['name']

        equip_val = data.get('equip_id')
        equip_val = str(equip_val) if equip_val else None

        c.execute("""INSERT INTO Usage_Logs (username, inventory_id, batch_code, quantity_used, unit_used, timestamp, activity_type, department, faculty, study, user_id, batch_id, department_id, faculty_id, study_id, remarks, balance_after, recorded_by, equip_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                  (user_name, data.get('inventory_id'), batch['batch_code'] or data.get('batch_code', ''), qty_used, data.get('unit', ''), datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), 'Lab Usage', dept_name, fac_name, std_name, user_id, data.get('batch_id'), dept_id, fac_id, std_id, data.get('remarks', ''), new_qty, session.get('user', 'System'), equip_val))
        
        conn.commit()
        conn.close()
        trigger_backup()
        return jsonify({"status": "success"})
        
    c.execute("SELECT * FROM Usage_Logs ORDER BY id DESC LIMIT 100")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return jsonify(rows)

@transactions_bp.route('/api/history', methods=['GET', 'POST'])
@login_required
def handle_history():
    conn = get_db()
    c = conn.cursor()
    if request.method == 'POST':
        data = request.get_json()
        c.execute("SELECT MAX(id) FROM History_Logs")
        max_id = c.fetchone()[0] or 0
        code = generate_semantic_id('HIS', max_id + 1, conn)
        c.execute("""INSERT INTO History_Logs (history_code, entity, entity_code, field_changed, old_value, new_value, changed_by, changed_date) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                  (code, data.get('entity', ''), data.get('entity_code', ''), data.get('field_changed', ''), data.get('old_value', ''), data.get('new_value', ''), session['user'], datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
        conn.commit(); conn.close(); trigger_backup()
        return jsonify({"status": "success"})
    c.execute("SELECT * FROM History_Logs ORDER BY id DESC")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return jsonify(rows)

@transactions_bp.route('/api/batches/template', methods=['GET'])
@admin_required
def batches_template():
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT item_code, material_name FROM Inventory_Master WHERE status='Active'")
    materials = c.fetchall()
    c.execute("SELECT id, name FROM Departments WHERE status='Active'")
    departments = c.fetchall()
    c.execute("SELECT id, name FROM Studies WHERE status='Active'")
    studies = c.fetchall()
    c.execute("SELECT id, vendor_name FROM Vendors")
    vendors = c.fetchall()
    conn.close()
    
    material_list = [f"{r['material_name'].replace(',', '')}_{r['item_code']}" for r in materials]
    dept_list = [f"{r['name'].replace(',', '')}_{r['id']}" for r in departments]
    study_list = [f"{r['name'].replace(',', '')}_{r['id']}" for r in studies]
    vendor_list = [f"{r['vendor_name'].replace(',', '')}_{r['id']}" for r in vendors]
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Batches Upload"
    
    def add_dropdown(data_list, title, col_letter):
        if not data_list: return
        ws_hidden = wb.create_sheet(title=title)
        ws_hidden.sheet_state = 'hidden'
        for idx, val in enumerate(data_list, start=1):
            ws_hidden.cell(row=idx, column=1, value=val)
        dv = DataValidation(type="list", formula1=f"'{title}'!$A$1:$A${len(data_list)}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{col_letter}2:{col_letter}1000')

    add_dropdown(material_list, "Material_List", "A")
    add_dropdown(dept_list, "Dept_List", "G")
    add_dropdown(study_list, "Study_List", "H")
    add_dropdown(vendor_list, "Vendor_List", "I")

    headers = ['InventoryCode', 'PONumber', 'LotNumber', 'ExpiryDate', 'DateFirstUsed', 'NumberOfPacks', 'DepartmentID', 'StudyID', 'VendorID', 'Status', 'Remarks']
    ws.append(headers)
    sample_mat = material_list[0] if material_list else "Sample_INV-000001"
    sample_dept = dept_list[0] if dept_list else "1"
    sample_study = study_list[0] if study_list else "1"
    sample_vendor = vendor_list[0] if vendor_list else "1"
    ws.append([sample_mat, 'PO-2026', 'LOT-999', '2026-12-31', '', 5, sample_dept, sample_study, sample_vendor, 'Active', 'Initial stock'])
    
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    response = make_response(out.read())
    response.headers['Content-Disposition'] = 'attachment; filename=Batches_Template.xlsx'
    response.headers['Content-type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response

@transactions_bp.route('/api/batches/upload', methods=['POST'])
@admin_required
def upload_batches():
    file = request.files.get('file')
    if not file: return jsonify({"status": "error", "message": "No file uploaded"}), 400
    try:
        wb = load_workbook(file, data_only=True)
        ws = wb.active
        rows = list(ws.rows)
        if len(rows) < 2: return jsonify({"status": "error", "message": "File is empty"}), 400
        headers = [str(cell.value).strip() if cell.value else f"Col{i}" for i, cell in enumerate(rows[0])]
        
        conn = get_db(); c = conn.cursor()
        c.execute("SELECT id, item_code, material_name, pack_qty, base_unit FROM Inventory_Master"); existing_inv = {r['item_code']: dict(r) for r in c.fetchall()}
        conn.close()

        parsed_data = []
        for row in rows[1:]:
            row_data = {headers[i]: cell.value for i, cell in enumerate(row)}
            inv_code = str(row_data.get('InventoryCode', '') or '').strip()
            if not inv_code or inv_code == 'None': continue
            
            extracted_code = inv_code.rpartition('_')[-1] if '_' in inv_code else inv_code
            
            inv_match = existing_inv.get(extracted_code)
            
            num_packs = float(row_data.get('NumberOfPacks') or 0)
            pack_qty = float(inv_match['pack_qty']) if inv_match and inv_match.get('pack_qty') else 1.0
            unit = str(inv_match['base_unit']) if inv_match and inv_match.get('base_unit') else 'Nos'
            quantity_received = num_packs * pack_qty
            
            dept_val = str(row_data.get('DepartmentID', '') or '').strip()
            dept_id = dept_val.rpartition('_')[-1] if '_' in dept_val else dept_val
            dept_id = int(dept_id) if dept_id.isdigit() else None

            study_val = str(row_data.get('StudyID', '') or '').strip()
            study_id = study_val.rpartition('_')[-1] if '_' in study_val else study_val
            study_id = int(study_id) if study_id.isdigit() else None

            vendor_val = str(row_data.get('VendorID', '') or '').strip()
            vendor_id = vendor_val.rpartition('_')[-1] if '_' in vendor_val else vendor_val
            vendor_id = int(vendor_id) if vendor_id.isdigit() else None
            
            parsed_data.append({
                "InventoryCode": extracted_code,
                "MatchedInventoryID": inv_match['id'] if inv_match else None,
                "MatchedMaterialName": inv_match['material_name'] if inv_match else "NOT FOUND",
                "PONumber": str(row_data.get('PONumber', '') or '').strip(),
                "LotNumber": str(row_data.get('LotNumber', '') or '').strip(),
                "ExpiryDate": str(row_data.get('ExpiryDate', '') or '').strip(),
                "DateFirstUsed": str(row_data.get('DateFirstUsed', '') or '').strip(),
                "NumberOfPacks": num_packs,
                "QuantityReceived": quantity_received,
                "Unit": unit,
                "DepartmentID": dept_id,
                "StudyID": study_id,
                "VendorID": vendor_id,
                "Status": str(row_data.get('Status', 'Active') or 'Active').strip(),
                "Remarks": str(row_data.get('Remarks', '') or '').strip()
            })
        return jsonify({"status": "success", "data": parsed_data})
    except Exception as e: return jsonify({"status": "error", "message": f"Failed to parse Excel: {str(e)}"}), 400

@transactions_bp.route('/api/batches/bulk', methods=['POST'])
@admin_required
def bulk_batches():
    data = request.get_json()
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT MAX(id) FROM Physical_Batches")
    max_id = c.fetchone()[0] or 0
    added_count = 0
    for row in data.get('rows', []):
        if not row.get('MatchedInventoryID'): continue
        if not row.get('StudyID'): continue
        
        max_id += 1
        try:
            code = generate_semantic_id('BAT', max_id, conn)
            qty = float(row.get('QuantityReceived', 0))
            c.execute("""INSERT INTO Physical_Batches (batch_code, inventory_id, po_number, lot_number, expiry_date, date_first_used, quantity_received, current_quantity, unit, department_id, study_id, vendor_id, status, remarks) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                      (code, row.get('MatchedInventoryID'), row.get('PONumber', ''), row.get('LotNumber', ''), row.get('ExpiryDate', ''), row.get('DateFirstUsed', ''), qty, qty, row.get('Unit', 'Nos'), row.get('DepartmentID'), row.get('StudyID'), row.get('VendorID'), row.get('Status', 'Active'), row.get('Remarks', '')))
            added_count += 1
        except sqlite3.IntegrityError: pass
    conn.commit()
    conn.close()
    trigger_backup()
    return jsonify({"status": "success", "message": f"Import complete: {added_count} added."})
